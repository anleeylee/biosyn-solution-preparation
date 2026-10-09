"""Export solution preparation results to XLSX / CSV."""

from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import List

from .calculator import SolutionResult

logger = logging.getLogger(__name__)


def export_solutions(results: List[SolutionResult], output_path: Path) -> None:
    """Export preparation results based on file extension."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    suffix = output_path.suffix.lower()

    if suffix == ".csv":
        _export_csv(results, output_path)
    else:
        _export_xlsx(results, output_path)
        csv_path = output_path.with_suffix(".csv")
        _export_csv(results, csv_path)

    logger.info("Exported %d solutions to %s", len(results), output_path)


def _export_csv(results: List[SolutionResult], path: Path) -> None:
    """Export flat CSV of all component results."""
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([
            "Solution", "Batch", "Final_Volume_mL", "Component",
            "Mode", "Amount", "Unit", "Warnings",
        ])
        for r in results:
            for cr in r.component_results:
                writer.writerow([
                    r.solution_name, r.batch, r.final_volume_ml,
                    cr.component_name, cr.mode, cr.amount, cr.amount_unit,
                    "; ".join(cr.warnings),
                ])


def _export_xlsx(results: List[SolutionResult], path: Path) -> None:
    """Export to multi-sheet XLSX."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    hdr = Font(bold=True)

    # Batch Summary
    ws = wb.active
    ws.title = "Batch Summary"
    ws.append(["Solution", "Batch", "Final Volume (mL)", "Components", "Warnings"])
    for c in ws[1]:
        c.font = hdr
    for r in results:
        ws.append([r.solution_name, r.batch, r.final_volume_ml,
                   len(r.component_results), len(r.warnings)])

    # Component Calculations
    ws2 = wb.create_sheet("Component Calculations")
    ws2.append(["Solution", "Component", "Mode", "Amount", "Unit", "Audit Trail", "Warnings"])
    for c in ws2[1]:
        c.font = hdr
    for r in results:
        for cr in r.component_results:
            ws2.append([
                r.solution_name, cr.component_name, cr.mode,
                cr.amount, cr.amount_unit, cr.audit_trail,
                "; ".join(cr.warnings),
            ])

    # Preparation Sheet (printable format)
    ws3 = wb.create_sheet("Preparation Sheet")
    row = 1
    for r in results:
        ws3.cell(row=row, column=1, value=f"Solution: {r.solution_name}").font = Font(bold=True, size=12)
        row += 1
        ws3.cell(row=row, column=1, value=f"Batch: {r.batch}")
        row += 1
        ws3.cell(row=row, column=1, value=f"Final Volume: {r.final_volume_ml} mL")
        row += 1
        ws3.cell(row=row, column=1, value="Component")
        ws3.cell(row=row, column=2, value="Amount").font = hdr
        row += 1
        for cr in r.component_results:
            ws3.cell(row=row, column=1, value=cr.component_name)
            ws3.cell(row=row, column=2, value=f"{cr.amount} {cr.amount_unit}")
            row += 1
        ws3.cell(row=row, column=1, value="Diluent")
        ws3.cell(row=row, column=2, value="q.s. to final volume")
        row += 2  # blank row between solutions

    # Warnings
    ws4 = wb.create_sheet("Warnings")
    ws4.append(["Solution", "Warning"])
    for c in ws4[1]:
        c.font = hdr
    for r in results:
        for w in r.warnings:
            ws4.append([r.solution_name, w])

    # Methods
    ws5 = wb.create_sheet("Methods")
    methods = [
        ("Method", "Description"),
        ("Molar", "mass = C × V × MW"),
        ("From Stock", "C1V1 = C2V2 → V1 = C2V2/C1"),
        ("Mass Conc.", "mass = C(g/L) × V(L)"),
        ("% w/v", "mass(g) = % / 100 × V(mL)"),
        ("% v/v", "volume(mL) = % / 100 × V(mL)"),
        ("Purity", "corrected mass = theoretical / purity_fraction"),
        ("", ""),
        ("Disclaimer", "q.s. / final volume instructions are planning notation. "
                       "Verify against lab SOP. Does not assess reagent compatibility, "
                       "chemical stability, pH adjustment, or storage stability."),
        ("Source", "BioSyn Solution Preparation Batch v1.0.0 — https://biosyn-inc.com/"),
    ]
    for row_data in methods:
        ws5.append(row_data)

    wb.save(path)
